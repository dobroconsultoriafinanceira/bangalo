# -*- coding: utf-8 -*-
"""DRE contábil: leitura do razão, importação e previsão."""
from datetime import date
from decimal import Decimal

import pytest
import xlrd  # noqa: F401  (garante a dependência instalada)
import xlwt

from app.extensions import db
from app.models.banco import DreContabilConta
from app.models.fluxo import Categoria, ConfigSistema, Lancamento
from app.services import dre_contabil as dc
from app.services import dre_previsao as prev

# (conta, classificação, nome, [(dia, histórico, contrapartida, débito, crédito)])
RAZAO = [
    ("552", "1.1.10.200.04", "BANCO ITAÚ", [(1, "VLR. REF. A PIX RECEBIDO", "504", 1000.0, 0.0)]),
    ("408", "4.1.10.100.3", "VENDA DE MERCADORIAS", [(31, "VLR FATURAMENTO NO MÊS", "504", 0.0, 100000.0)]),
    ("417", "4.1.20.100.03", "(-) DEVOLUÇÃO DE VENDA", [(1, "DEVOLUÇÃO NESTA DATA", "504", 10000.0, 0.0)]),
    ("480", "4.1.20.300.08", "(-) SIMPLES NACIONAL", [(31, "SIMPLES", "479", 8000.0, 0.0)]),
    ("426", "4.1.20.300.02", "(-) ICMS", [(31, "ICMS", "172", 2000.0, 0.0)]),
    ("712", "3.1.10.100.01", "CUSTO DAS MERCADORIAS VENDIDAS", [(31, "CMV DO MES", "853", 30000.0, 0.0)]),
    ("331", "3.2.20.100.01", "SALÁRIOS E ORDENADOS", [(31, "FOLHA", "187", 12000.0, 0.0)]),
    ("332", "3.2.20.100.02", "PRÓ-LABORE", [(31, "PRO LABORE", "188", 8475.55, 0.0)]),
    ("357", "3.2.20.400.04", "ALUGUEL DE IMÓVEL", [(10, "ALUGUEL", "586", 7000.0, 0.0)]),
    ("1791", "3.2.20.400.54", "DESPESA COM MÚSICO", [(5, "PG. NF MUSICO", "552", 3000.0, 0.0)]),
    ("348", "3.2.20.300.03", "IPTU", [(20, "IPTU", "552", 1000.0, 0.0)]),
    ("1049", "3.2.20.500.11", "DESPESAS COM ADM. DE CARTÃO", [(31, "TAXA CARTAO", "504", 1200.0, 0.0)]),
    ("372", "3.2.20.500.05", "JUROS", [(27, "JUROS", "552", 800.0, 0.0)]),
    ("1652", "3.2.21.600.01", "DEPRECIAÇÃO DE INSTALAÇÕES", [(31, "DEPRECIACAO", "591", 500.0, 0.0)]),
    ("896", "4.1.30.100.06", "RENDIMENTO DE APLICAÇÃO", [(31, "RENDIMENTO", "1025", 0.0, 1500.0)]),
    ("515", "4.1.50.100.06", "BONIFICAÇÕES E BRINDES", [(15, "BONIFICACAO", "853", 0.0, 300.0)]),
]


@pytest.fixture()
def razao(tmp_path):
    """Gera um razão .xls no mesmo layout do arquivo da contabilidade."""
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Razão")
    ws.write(0, 0, "Empresa:")
    ws.write(6, 0, "Data")
    linha = 7
    for conta, classificacao, nome, lancamentos in RAZAO:
        ws.write(linha, 0, "Conta:")
        ws.write(linha, 1, conta)
        ws.write(linha, 2, classificacao)
        ws.write(linha, 5, nome)
        linha += 1
        ws.write(linha, 2, "SALDO ANTERIOR")  # linha sem data: deve ser ignorada
        linha += 1
        total_d = total_c = 0.0
        for dia, historico, contrapartida, debito, credito in lancamentos:
            serial = (date(2026, 7, dia) - date(1899, 12, 30)).days
            ws.write(linha, 0, serial)
            ws.write(linha, 2, historico)
            ws.write(linha, 7, contrapartida)
            if debito:
                ws.write(linha, 8, debito)
            if credito:
                ws.write(linha, 9, credito)
            total_d += debito
            total_c += credito
            linha += 1
        # subtotal repetido pelo arquivo, sem data — não pode entrar na conta
        ws.write(linha, 8, total_d)
        ws.write(linha, 9, total_c)
        linha += 2
    caminho = tmp_path / "razao.xls"
    wb.save(caminho)
    return caminho


def test_le_razao_e_monta_o_dre(razao):
    lancamentos = dc.ler_razao(razao)
    assert len(lancamentos) == sum(len(l[3]) for l in RAZAO)
    assert dc.periodo(lancamentos) == (date(2026, 7, 1), date(2026, 7, 31))

    dre = dc.dre_do_razao(lancamentos)
    assert dre["receita_bruta"] == Decimal("100000.00")
    assert dre["linhas"]["deducoes"].valor == Decimal("20000.00")   # gorjeta + Simples + ICMS
    assert dre["receita_liquida"] == Decimal("80000.00")
    assert dre["linhas"]["cmv"].valor == Decimal("30000.00")
    assert dre["lucro_bruto"] == Decimal("50000.00")
    assert dre["linhas"]["pessoal"].valor == Decimal("20475.55")
    assert dre["linhas"]["gerais"].valor == Decimal("10000.00")
    assert dre["linhas"]["financeiras"].valor == Decimal("2000.00")
    assert dre["linhas"]["depreciacao"].valor == Decimal("500.00")
    # 50.000 − (20.475,55 + 10.000 + 1.000 + 2.000 + 500) = 16.024,45
    assert dre["resultado_operacional"] == Decimal("16024.45")
    assert dre["resultado"] == Decimal("17824.45")  # + rendimento 1.500 + bonificações 300
    assert len(dc.lancamentos_do_banco(lancamentos)) == 1


def test_importa_razao_e_remonta_do_banco(app, razao):
    rel = dc.importar_razao(razao)
    db.session.commit()
    assert (rel["ano"], rel["mes"]) == (2026, 7)
    assert rel["receita_bruta"] == Decimal("100000.00")
    assert db.session.query(DreContabilConta).count() == rel["contas"]
    assert dc.meses_importados() == [(2026, 7)]

    salvo = dc.dre_importado(2026, 7)
    assert salvo["resultado"] == Decimal("17824.45")
    assert salvo["linhas"]["gerais"].contas[0]["nome"] == "ALUGUEL DE IMÓVEL"

    # reimportar o mesmo mês substitui, não duplica
    dc.importar_razao(razao)
    db.session.commit()
    assert db.session.query(DreContabilConta).count() == rel["contas"]
    assert dc.dre_importado(2026, 8) is None


def _nosso_fluxo_de_julho():
    cats = {nome: Categoria(nome=nome, tipo=tipo, grupo=grupo) for nome, tipo, grupo in [
        ("Visa Crédito", "entrada", "Vendas - Repasse Stone"),
        ("Compras", "saida", "Compras"),
        ("Salários", "saida", "Folha/Salários"),
        ("Aluguel", "saida", "Despesas Fixas"),
        ("IPTU", "saida", "Despesas Fixas"),
        ("Músicos", "saida", "Demais Salários"),
    ]}
    db.session.add_all(cats.values())
    db.session.flush()
    for dia, nome, valor in [(5, "Visa Crédito", "90000"), (10, "Compras", "25000"),
                             (6, "Salários", "12000"), (10, "Aluguel", "7000"),
                             (20, "IPTU", "1000"), (5, "Músicos", "3000")]:
        db.session.add(Lancamento(data=date(2026, 7, dia), categoria_id=cats[nome].id, valor=Decimal(valor)))
    db.session.commit()


def test_calibra_com_o_razao_e_preve_o_mes(app, razao):
    _nosso_fluxo_de_julho()
    dc.importar_razao(razao)
    db.session.commit()

    params = prev.calibrar(2026, 7)
    db.session.commit()
    assert params["calibrado_com"] == "07/2026"
    assert Decimal(params["pct_gorjeta"]) == Decimal("0.10000")      # 10.000 / 100.000
    assert Decimal(params["pct_cmv"]) == Decimal("0.30000")
    assert Decimal(params["fator_gerais"]) == Decimal("1.0000")      # 10.000 / (7.000 + 3.000)

    # com o faturamento informado, a previsão do próprio julho reproduz o razão
    prev.definir_faturamento(2026, 7, Decimal("100000.00"))
    db.session.commit()
    previsto = prev.prever(2026, 7)
    assert previsto["receita_bruta"] == Decimal("100000.00")
    assert previsto["linhas"]["deducoes"].valor == Decimal("20000.00")
    assert previsto["linhas"]["pessoal"].valor == Decimal("20475.55")
    assert previsto["resultado"] == Decimal("17824.45")
    assert "informado" in previsto["faturamento_origem"]

    comparacao = prev.comparar(2026, 7)
    assert all(l["diferenca"] == Decimal("0.00") for l in comparacao["linhas"])
    assert comparacao["totais"][-1]["real"] == Decimal("17824.45")


def test_tela_dre_contabil(app, client, razao):
    from app.models.usuario import Usuario

    _nosso_fluxo_de_julho()
    dc.importar_razao(razao)
    prev.calibrar(2026, 7)
    db.session.commit()
    u = Usuario(nome="C", email="c@x.com", role="consultoria")
    u.definir_senha("senha1234")
    db.session.add(u)
    db.session.commit()
    client.post("/login", data={"email": "c@x.com", "senha": "senha1234"})

    resp = client.get("/dre/contabil?ano=2026&mes=7")
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200
    for trecho in ("DRE contábil", "Prévia interna", "Contabilidade", "ALUGUEL DE IMÓVEL", "Lucro líquido"):
        assert trecho in html

    # mês sem razão importado: mostra só a previsão
    resp = client.get("/dre/contabil?ano=2026&mes=8")
    assert resp.status_code == 200
    assert "ainda não enviou" in resp.get_data(as_text=True)

    # informar o faturamento do mês
    resp = client.post("/dre/contabil/faturamento",
                       data={"ano": 2026, "mes": 8, "faturamento": "100.000,00"})
    assert resp.status_code in (302, 303)
    assert prev.faturamento_informado(2026, 8) == Decimal("100000.00")


def test_previsao_sem_faturamento_informado_estima_pelo_caixa(app):
    _nosso_fluxo_de_julho()
    previsto = prev.prever(2026, 7)
    assert previsto["receita_bruta"] > Decimal("90000.00")  # bruto = recebido + taxa
    assert "estimado" in previsto["faturamento_origem"]
    assert prev.comparar(2026, 7)["real"] is None
    assert ConfigSistema.obter(prev.CHAVE_PARAMETROS) is None  # previsão não altera parâmetros


# ── Fechamento Clareza: prévia, documentos e razão detalhado ─────────────────

def _entrar(client, role="consultoria"):
    from app.models.usuario import Usuario

    u = Usuario(nome="Pessoa", email=f"{role}@x.com", role=role)
    u.definir_senha("senha1234")
    db.session.add(u)
    db.session.commit()
    client.post("/login", data={"email": f"{role}@x.com", "senha": "senha1234"})


def test_previa_nao_grava_e_importacao_guarda_documento_e_lancamentos(app, client, razao, tmp_path):
    from app.models.documentos import DocumentoContabil, RazaoLancamento

    app.instance_path = str(tmp_path / "instance")
    _entrar(client)
    with open(razao, "rb") as fh:
        resp = client.post("/dre/contabil/razao/previa", data={"ano": 2026, "mes": 7, "arquivo": (fh, "razao julho.xls")},
                           content_type="multipart/form-data")
    assert resp.status_code in (302, 303)
    local = resp.headers["Location"]
    assert "previa=" in local
    assert db.session.query(DreContabilConta).count() == 0  # prévia não grava

    html = client.get(local).get_data(as_text=True)
    assert "Confirmar importação do razão" in html and "Julho 2026" in html
    token = local.split("previa=")[1].split("&")[0]

    resp = client.post("/dre/contabil/razao", data={"ano": 2026, "mes": 7, "previa": token, "nome": "razao julho.xls"})
    assert resp.status_code in (302, 303)
    assert db.session.query(DreContabilConta).filter_by(ano=2026, mes=7).count() > 0
    assert db.session.query(RazaoLancamento).filter_by(ano=2026, mes=7, conta="357").count() == 1
    doc = db.session.execute(db.select(DocumentoContabil)).scalar_one()
    assert (doc.tipo, doc.nome_arquivo, doc.situacao) == ("razao", "razao julho.xls", "recebido")
    assert client.get(f"/dre/contabil/documentos/{doc.id}/arquivo").status_code == 200

    # segunda prévia do mesmo mês avisa a substituição
    with open(razao, "rb") as fh:
        local = client.post("/dre/contabil/razao/previa", data={"ano": 2026, "mes": 7, "arquivo": (fh, "razao.xls")},
                            content_type="multipart/form-data").headers["Location"]
    assert "Substituir razão do mês" in client.get(local).get_data(as_text=True)

    html = client.get("/dre/contabil?ano=2026&mes=7&aba=razao&conta=357").get_data(as_text=True)
    assert "ALUGUEL DE IMÓVEL" in html and "Total da conta" in html


def test_documento_conferido_pela_gerencia_e_aprovado_pela_consultoria(app, client, tmp_path):
    from io import BytesIO

    from app.models.documentos import DocumentoContabil

    app.instance_path = str(tmp_path / "instance")
    _entrar(client, "gerencia")
    resp = client.post("/dre/contabil/documentos",
                       data={"ano": 2026, "mes": 8, "tipo": "dre_pdf", "arquivo": (BytesIO(b"%PDF-1.4"), "DRE agosto.pdf")},
                       content_type="multipart/form-data")
    assert resp.status_code in (302, 303)
    doc = db.session.execute(db.select(DocumentoContabil)).scalar_one()

    # formato errado é recusado
    client.post("/dre/contabil/documentos", data={"ano": 2026, "mes": 8, "tipo": "dre_pdf",
                "arquivo": (BytesIO(b"x"), "planilha.exe")}, content_type="multipart/form-data")
    assert db.session.query(DocumentoContabil).count() == 1

    client.post(f"/dre/contabil/documentos/{doc.id}/situacao", data={"situacao": "conferido"})
    assert db.session.get(DocumentoContabil, doc.id).situacao == "conferido"
    assert client.post(f"/dre/contabil/documentos/{doc.id}/situacao", data={"situacao": "aprovado"}).status_code == 403


def _converter_para_xlsx(origem_xls, destino):
    """Mesmo conteúdo do razão, salvo em .xlsx (como sai do Excel ao recuperar um .xls)."""
    import openpyxl
    import xlrd

    ws_origem = xlrd.open_workbook(origem_xls).sheet_by_index(0)
    wb = openpyxl.Workbook()
    ws = wb.active
    for l in range(ws_origem.nrows):
        for col in range(ws_origem.ncols):
            valor = ws_origem.cell_value(l, col)
            if valor not in ("", None):
                ws.cell(row=l + 1, column=col + 1, value=valor)
    wb.save(destino)
    return destino


def test_razao_em_xlsx_da_no_mesmo_resultado(razao, tmp_path):
    xlsx = _converter_para_xlsx(razao, tmp_path / "razao.xlsx")
    do_xls, do_xlsx = dc.dre_do_razao(dc.ler_razao(razao)), dc.dre_do_razao(dc.ler_razao(xlsx))
    assert do_xlsx["receita_bruta"] == do_xls["receita_bruta"]
    assert do_xlsx["resultado"] == do_xls["resultado"]
    assert dc.periodo(dc.ler_razao(xlsx)) == dc.periodo(dc.ler_razao(razao))


def test_razao_corrompido_avisa_em_vez_de_quebrar(razao, tmp_path):
    truncado = tmp_path / "razao_truncado.xls"
    truncado.write_bytes(razao.read_bytes()[:-1024])  # arquivo que chegou incompleto
    with pytest.raises(dc.RazaoIlegivel, match="incompleto ou corrompido"):
        dc.ler_razao(truncado)


def test_tela_aceita_razao_xlsx(app, client, razao, tmp_path):
    _entrar(client)
    xlsx = _converter_para_xlsx(razao, tmp_path / "razao agosto.xlsx")
    with open(xlsx, "rb") as fh:
        resp = client.post("/dre/contabil/razao/previa", data={"ano": 2026, "mes": 7, "arquivo": (fh, "razao agosto.xlsx")},
                           content_type="multipart/form-data")
    assert "previa=" in resp.headers["Location"] and "ext=.xlsx" in resp.headers["Location"]
    token = resp.headers["Location"].split("previa=")[1].split("&")[0]
    resp = client.post("/dre/contabil/razao", data={"ano": 2026, "mes": 7, "previa": token, "nome": "razao agosto.xlsx"})
    assert resp.status_code in (302, 303)
    from app.models.documentos import DocumentoContabil

    doc = db.session.execute(db.select(DocumentoContabil).filter_by(tipo="razao")).scalar_one()
    assert doc.nome_arquivo == "razao agosto.xlsx"
    assert db.session.query(DreContabilConta).filter_by(ano=2026, mes=7).count() > 0
