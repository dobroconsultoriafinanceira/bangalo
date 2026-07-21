# -*- coding: utf-8 -*-
"""Importa a planilha 'Metas de Faturamento 2026'.

Abas usadas:
  - 'Faturamento Histórico': matriz Mês × Ano (2022–2026) → faturamento_historico
  - 'Registro Diário 2026':  Data | Dia | Aberto? | Faturamento → faturamento_diario
  - 'Metas':                 premissas (pesos, crescimento) → premissa_meta

Idempotente: usa upsert por (ano, mes) / data. Emite relatório de conferência.
"""
from decimal import Decimal
from pathlib import Path

import openpyxl

from app.extensions import db
from app.importers.comum import para_data, para_decimal
from app.models.metas import FaturamentoDiario, FaturamentoHistorico, PremissaMeta

MESES_ORDEM = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN",
               "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]


def importar(caminho: Path) -> list[str]:
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    relatorio: list[str] = []
    relatorio += _importar_historico(wb)
    relatorio += _importar_registro_diario(wb)
    relatorio += _importar_premissas(wb)
    wb.close()
    db.session.flush()
    return relatorio


def _importar_historico(wb) -> list[str]:
    if "Faturamento Histórico" not in wb.sheetnames:
        return ["Aba 'Faturamento Histórico' não encontrada — pulada."]
    ws = wb["Faturamento Histórico"]
    linhas = list(ws.iter_rows(values_only=True))
    # linha 2 (índice 1): Mês | 2022 | 2023 | 2024 | 2025 | 2026
    header = linhas[1]
    anos = [int(a) for a in header[1:] if isinstance(a, (int, float))]
    total_importado = 0
    total_valor = Decimal("0")
    for row in linhas[2:]:
        if not row or not row[0]:
            continue
        rotulo = str(row[0]).strip().upper()
        if rotulo not in MESES_ORDEM:
            continue
        mes = MESES_ORDEM.index(rotulo) + 1
        for i, ano in enumerate(anos):
            valor = para_decimal(row[1 + i])
            if valor is None:
                continue
            _upsert_historico(ano, mes, valor)
            total_importado += 1
            total_valor += valor
    return [f"Faturamento histórico: {total_importado} valores importados (soma R$ {total_valor})."]


def _upsert_historico(ano: int, mes: int, valor: Decimal) -> None:
    reg = db.session.execute(
        db.select(FaturamentoHistorico).filter_by(ano=ano, mes=mes)
    ).scalar_one_or_none()
    if reg:
        reg.valor = valor
    else:
        db.session.add(FaturamentoHistorico(ano=ano, mes=mes, valor=valor))


def _importar_registro_diario(wb) -> list[str]:
    if "Registro Diário 2026" not in wb.sheetnames:
        return ["Aba 'Registro Diário 2026' não encontrada — pulada."]
    ws = wb["Registro Diário 2026"]
    importados = 0
    com_valor = 0
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        data = para_data(row[0])
        if not data:
            continue
        aberto = str(row[2]).strip().lower().startswith("aberto") if len(row) > 2 and row[2] else True
        faturamento = para_decimal(row[3]) if len(row) > 3 else None
        _upsert_diario(data, aberto, faturamento)
        importados += 1
        if faturamento is not None:
            com_valor += 1
    return [f"Registro diário: {importados} dias importados ({com_valor} com faturamento)."]


def _upsert_diario(data, aberto, faturamento) -> None:
    reg = db.session.execute(
        db.select(FaturamentoDiario).filter_by(data=data)
    ).scalar_one_or_none()
    if reg:
        reg.aberto = aberto
        if faturamento is not None:
            reg.faturamento = faturamento
    else:
        db.session.add(FaturamentoDiario(data=data, aberto=aberto, faturamento=faturamento))


def _importar_premissas(wb) -> list[str]:
    if "Metas" not in wb.sheetnames:
        return ["Aba 'Metas' não encontrada — premissas mantidas no padrão."]
    ws = wb["Metas"]
    valores = {}
    for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        if not row or not row[0]:
            continue
        rotulo = str(row[0]).strip().lower()
        num = para_decimal(row[1]) if len(row) > 1 else None
        if "mesmo mês" in rotulo and num is not None:
            valores["peso_mes_anterior"] = num
        elif "média histórica" in rotulo and num is not None:
            valores["peso_media_historica"] = num
        elif "crescimento" in rotulo and num is not None:
            valores["crescimento_alvo"] = num

    ano = 2026
    reg = db.session.execute(db.select(PremissaMeta).filter_by(ano=ano)).scalar_one_or_none()
    if not reg:
        reg = PremissaMeta(ano=ano)
        db.session.add(reg)
    if "peso_mes_anterior" in valores:
        reg.peso_mes_anterior = valores["peso_mes_anterior"]
    if "peso_media_historica" in valores:
        reg.peso_media_historica = valores["peso_media_historica"]
    if "crescimento_alvo" in valores:
        reg.crescimento_alvo = valores["crescimento_alvo"]
    return [f"Premissas {ano}: peso_ant={reg.peso_mes_anterior} peso_hist={reg.peso_media_historica} cresc={reg.crescimento_alvo}."]
